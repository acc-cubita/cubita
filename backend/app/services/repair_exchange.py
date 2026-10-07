"""Scoped XLSX exchange. Imports are historical annotations, never transactions."""
from io import BytesIO
from zipfile import ZipFile,BadZipFile
from hashlib import sha256
from datetime import date,datetime
from fastapi import HTTPException
from openpyxl import Workbook,load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from app.models.repair_completion import RepairHistoryImport,RepairHistoricalRecord
from app.services import repair as core

HEADERS=['external_id','date','customer','device','serial','issue','result']
MAX_ROWS=2000


def text(value):
    value='' if value is None else str(value)
    return value


def workbook(rows,headers):
    book=Workbook();sheet=book.active;sheet.title='تعمیرگاه';sheet.sheet_view.rightToLeft=True
    for values in [headers,*rows]:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str): cell.data_type='s'
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for column in sheet.columns: sheet.column_dimensions[column[0].column_letter].width=24
    stream=BytesIO();book.save(stream);return stream.getvalue()


def preview(db,p,branch_id,source_id,content):
    core.branch(db,p,branch_id)
    source_id=source_id.strip()
    if not source_id or len(source_id)>120: raise HTTPException(422,'شناسهٔ منبع لازم است و حداکثر ۱۲۰ نویسه دارد.')
    if len(content)>5*1024*1024: raise HTTPException(413,'فایل حداکثر پنج مگابایت است.')
    digest=sha256(content).hexdigest()
    db.query(core.RepairBranch).filter_by(id=branch_id).with_for_update().one()
    existing=db.query(RepairHistoryImport).filter_by(branch_id=branch_id,source_id=source_id,digest=digest).first()
    if existing:
        if existing.recorded_by_id!=p.user.id and not p.has_permission('repair','approve'): raise HTTPException(403,'پیش‌نمایش متعلق به کارمند دیگری است.')
        return existing
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos=archive.infolist()
            if len(infos)>200 or sum(i.file_size for i in infos)>20*1024*1024: raise ValueError('اندازهٔ بازشدهٔ فایل بیش از حد است')
            if any('vbaProject' in i.filename or 'externalLinks/' in i.filename for i in infos): raise ValueError('ماکرو و پیوند خارجی مجاز نیست')
        book=load_workbook(BytesIO(content),read_only=True,data_only=False,keep_links=False)
        if len(book.worksheets)!=1: raise ValueError('فایل باید یک برگه داشته باشد')
        sheet=book.active
        if sheet.max_row is None or sheet.max_column is None or sheet.max_row>MAX_ROWS+1 or sheet.max_column!=len(HEADERS): raise ValueError('حداکثر ۲۰۰۰ ردیف و هفت ستون مجاز است')
        iterator=sheet.iter_rows();header=next(iterator)
        if [text(c.value).strip() for c in header]!=HEADERS: raise ValueError('ستون‌ها باید مطابق فایل نمونه باشند')
        rows=[];seen=set()
        for number,cells in enumerate(iterator,2):
            if number>MAX_ROWS+1: raise ValueError('حداکثر ۲۰۰۰ ردیف مجاز است')
            if all(c.value is None for c in cells): continue
            errors=[];data={key:text(cell.value).strip() for key,cell in zip(HEADERS,cells)}
            if any(c.data_type=='f' for c in cells): errors.append('فرمول مجاز نیست')
            if any(len(v)>2000 for v in data.values()): errors.append('متن بیش از حد طولانی است')
            if not data['external_id'] or len(data['external_id'])>120: errors.append('شناسهٔ سابقه لازم است و حداکثر ۱۲۰ نویسه دارد')
            if not data['device'] or not data['customer']: errors.append('نام مشتری و دستگاه لازم است')
            try:
                raw=cells[1].value
                data['date']=(raw.date() if isinstance(raw,datetime) else raw if isinstance(raw,date) else date.fromisoformat(data['date'])).isoformat()
            except (ValueError,TypeError): errors.append('تاریخ میلادی معتبر به صورت YYYY-MM-DD لازم است')
            duplicate=data['external_id'] in seen or db.query(RepairHistoricalRecord.id).filter_by(branch_id=branch_id,source_id=source_id,external_id=data['external_id']).first() is not None
            seen.add(data['external_id']);rows.append({'row':number,'data':data,'errors':errors,'duplicate':duplicate})
        book.close()
        if not rows: raise ValueError('فایل سابقه‌ای ندارد')
    except (ValueError,BadZipFile,KeyError,StopIteration,InvalidFileException,ParseError,DefusedXmlException) as exc:
        raise HTTPException(422,'فایل Excel معتبر نیست: '+str(exc)) from exc
    row=RepairHistoryImport(branch_id=branch_id,source_id=source_id,digest=digest,rows=rows,recorded_by_id=p.user.id)
    db.add(row);db.flush();return row


def import_row(db,p,import_id,lock=False):
    query=db.query(RepairHistoryImport).filter_by(id=import_id)
    row=(query.with_for_update() if lock else query).one_or_none()
    if row is None: raise HTTPException(404,'پیش‌نمایش پیدا نشد.')
    core.branch(db,p,row.branch_id)
    if row.recorded_by_id!=p.user.id and not p.has_permission('repair','approve'): raise HTTPException(403,'پیش‌نمایش متعلق به کارمند دیگری است.')
    return row


def commit(db,p,row,version):
    core.check_version(row,version)
    if row.status!='preview': raise HTTPException(409,'این ورود قبلاً نهایی شده است.')
    # Branch lock serializes imports of the same source across different uploads.
    db.query(core.RepairBranch).filter_by(id=row.branch_id).with_for_update().one()
    rows=[]
    for entry in row.rows:
        entry={**entry}
        if entry['errors']: entry['outcome']='invalid'
        elif db.query(RepairHistoricalRecord.id).filter_by(branch_id=row.branch_id,source_id=row.source_id,external_id=entry['data']['external_id']).first(): entry['outcome']='duplicate'
        else:
            db.add(RepairHistoricalRecord(branch_id=row.branch_id,import_id=row.id,source_id=row.source_id,external_id=entry['data']['external_id'],data=entry['data']));db.flush();entry['outcome']='imported'
        rows.append(entry)
    row.rows=rows;row.status='committed';row.version+=1;db.flush();return row
