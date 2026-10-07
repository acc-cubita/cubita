from io import BytesIO
from uuid import uuid4
from openpyxl import load_workbook,Workbook
from tests.test_repair import intake,admit
from app.services.repair_exchange import workbook,HEADERS
from app.models.repair_completion import RepairHistoricalRecord
from app.models.repair import RepairCase,RepairDevice
from app.models.inventory import StockLedger,Contact
from app.models.accounting import JournalEntry
from app.models.invoices import PurchaseInvoice,SalesInvoice


def counts(db):
    return tuple(db.query(m).count() for m in [RepairCase,RepairDevice,Contact,StockLedger,JournalEntry,PurchaseInvoice,SalesInvoice])


def test_history_import_preview_errors_dedup_only_history(client,db,intake):
    before=counts(db)
    rows=[['old-1','2025-01-01','مشتری','دستگاه','S1','ایراد','رفع شد'],['old-1','2025-01-01','مشتری','دستگاه','','',''],['bad','bad-date','مشتری','دستگاه','','','']]
    content=workbook(rows,HEADERS)
    path='/api/repair/branches/'+intake['branch_id']+'/history-imports'
    preview=client.post(path,data={'source_id':'old-workshop'},files={'file':('history.xlsx',content)})
    assert preview.status_code==201,preview.text
    data=preview.json();assert data['rows'][1]['duplicate'] and data['rows'][2]['errors']
    assert counts(db)==before and db.query(RepairHistoricalRecord).count()==0
    assert client.post(path,data={'source_id':'old-workshop'},files={'file':('history.xlsx',content)}).json()['id']==data['id']
    key=str(uuid4());commit='/api/repair/history-imports/'+data['id']+'/commit'
    result=client.post(commit,json={'version':1},headers={'Idempotency-Key':key})
    assert result.status_code==200,result.text
    assert [r['outcome'] for r in result.json()['rows']]==['imported','duplicate','invalid']
    assert client.post(commit,json={'version':1},headers={'Idempotency-Key':key}).json()['id']==data['id']
    assert counts(db)==before and db.query(RepairHistoricalRecord).count()==1
    changed=workbook([rows[0]],HEADERS)
    again=client.post(path,data={'source_id':'old-workshop'},files={'file':('other.xlsx',changed)})
    assert again.json()['rows'][0]['duplicate']


def test_excel_export_filters_literal_cells_no_credentials(client,db,intake):
    row=admit(client,intake).json()
    exported=client.get('/api/repair/exchange.xlsx',params={'branch_id':intake['branch_id'],'status':'accepted','financial':True})
    assert exported.status_code==200,exported.text
    book=load_workbook(BytesIO(exported.content));values=list(book.active.values)
    assert len(values)==2 and values[1][0]==row['id']
    assert not any('رمز' in str(c) or 'توکن' in str(c) for r in values for c in r)
    assert len(list(load_workbook(BytesIO(client.get('/api/repair/exchange.xlsx',params={'status':'closed'}).content)).active.values))==1
    literal=load_workbook(BytesIO(workbook([['=HYPERLINK("evil")']],['test']))).active['A2']
    assert literal.data_type=='s'


def test_history_formula_is_row_error_invalid_file(client,intake):
    book=Workbook();book.active.append(HEADERS);book.active.append(['ID','2025-01-01','=1+1','دستگاه','','',''])
    stream=BytesIO();book.save(stream)
    path='/api/repair/branches/'+intake['branch_id']+'/history-imports'
    result=client.post(path,data={'source_id':'formula'},files={'file':('test.xlsx',stream.getvalue())})
    assert result.status_code==201,result.text
    assert 'فرمول مجاز نیست' in result.json()['rows'][0]['errors']
    assert client.post(path,data={'source_id':'bad'},files={'file':('test.xlsx',b'not-excel')}).status_code==422
